"""Extract original BACE endpoint provenance from the publisher's SI archive.

Read-only with respect to the source ZIP and the project's raw BACE CSV.
Run with a Python environment containing openpyxl and pandas.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ZIP = ROOT / "data/external_v2/original_bace_source/ci6b00290_si_001.zip"
RAW_BACE = ROOT / "data/raw/bace.csv"
PROVENANCE = ROOT / "data/external_v2/original_bace_source/source_provenance.csv"
LABEL_AUDIT = ROOT / "results_v2/external_wp2_original_label_audit.csv"
SUMMARY = ROOT / "results_v2/external_wp2_original_label_summary.json"
EXPECTED_MD5 = "aa483ebebcb0b18fa5d639700db32f60"


def digest(path: Path, algorithm: str) -> str:
    h = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    if digest(SOURCE_ZIP, "md5") != EXPECTED_MD5:
        raise ValueError("The original BACE SI archive does not match Figshare's MD5")
    with zipfile.ZipFile(SOURCE_ZIP) as archive:
        workbook = load_workbook(
            io.BytesIO(archive.read("bace1547_sin.xlsx")),
            read_only=True,
            data_only=True,
        )
        sheet = workbook["S1-bace1547_liter"]
        row_iter = sheet.iter_rows(values_only=True)
        headers = list(next(row_iter))
        source_rows = []
        for excel_row, cells in enumerate(row_iter, start=2):
            record = dict(zip(headers, cells))
            source_rows.append(
                {
                    "source_excel_row": excel_row,
                    "CID": record.get("CID"),
                    "source_smiles": record.get("Canonical_Smiles"),
                    "source_endpoint": record.get("STANDARD_TYPE"),
                    "source_relation": record.get("RELATION"),
                    "source_value": record.get("STANDARD_VALUE"),
                    "source_units": record.get("STANDARD_UNITS"),
                    "source_pIC50": record.get("pIC50"),
                    "source_class": record.get("Class"),
                    "source_pubmed_id": record.get("PUBMED_ID"),
                    "source_journal": record.get("JOURNAL"),
                    "source_year": record.get("YEAR"),
                    "source_company": record.get("Company"),
                    "source_compound_key": record.get("COMPOUND_KEY"),
                    "source_parent_chembl_id": record.get("PARENT_CMPD_CHEMBLID"),
                }
            )
        workbook.close()

    PROVENANCE.parent.mkdir(parents=True, exist_ok=True)
    with PROVENANCE.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(source_rows[0]))
        writer.writeheader()
        writer.writerows(source_rows)

    by_cid = {str(row["CID"]): row for row in source_rows if row["CID"] is not None}
    raw = pd.read_csv(RAW_BACE, usecols=["CID", "Class", "pIC50"])
    audited = []
    for record in raw.itertuples(index=False):
        source = by_cid.get(str(record.CID))
        source_endpoint = source["source_endpoint"] if source else None
        threshold_label = int(float(record.pIC50) >= 7.0)
        audited.append(
            {
                "CID": record.CID,
                "deepchem_class": int(record.Class),
                "deepchem_pIC50": record.pIC50,
                "pIC50_ge_7": threshold_label,
                "class_threshold_disagreement": int(int(record.Class) != threshold_label),
                "source_match_by_CID": int(source is not None),
                "source_endpoint": source_endpoint,
                "source_class": source["source_class"] if source else None,
                "source_pubmed_id": source["source_pubmed_id"] if source else None,
                "source_year": source["source_year"] if source else None,
                "source_company": source["source_company"] if source else None,
                "source_excel_row": source["source_excel_row"] if source else None,
            }
        )
    LABEL_AUDIT.parent.mkdir(parents=True, exist_ok=True)
    with LABEL_AUDIT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(audited[0]))
        writer.writeheader()
        writer.writerows(audited)

    original_ic50 = [
        row for row in source_rows
        if row["source_endpoint"] == "IC50"
        and row["source_class"] in (0, 1)
        and row["source_pIC50"] is not None
    ]
    summary = {
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_url": "https://acs.figshare.com/articles/dataset/Computational_Modeling_of_Secretase_1_BACE-1_Inhibitors_Using_Ligand_Based_Approaches/4003788",
        "source_figshare_file_id": 6431988,
        "source_zip_sha256": digest(SOURCE_ZIP, "sha256"),
        "original_si_rows": len(source_rows),
        "original_si_endpoint_counts": dict(Counter(str(row["source_endpoint"]) for row in source_rows)),
        "original_si_ic50_with_class": len(original_ic50),
        "original_si_ic50_threshold_disagreements": sum(
            int((row["source_class"] == 1) != (float(row["source_pIC50"]) >= 7.0))
            for row in original_ic50
        ),
        "deepchem_rows": len(audited),
        "deepchem_rows_missing_original_CID": sum(not row["source_match_by_CID"] for row in audited),
        "deepchem_class_threshold_disagreements": sum(row["class_threshold_disagreement"] for row in audited),
        "deepchem_disagreement_endpoints": dict(Counter(
            str(row["source_endpoint"])
            for row in audited if row["class_threshold_disagreement"]
        )),
        "source_provenance_sha256": digest(PROVENANCE, "sha256"),
        "label_audit_sha256": digest(LABEL_AUDIT, "sha256"),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({k: summary[k] for k in (
        "original_si_ic50_with_class",
        "original_si_ic50_threshold_disagreements",
        "deepchem_class_threshold_disagreements",
        "deepchem_rows_missing_original_CID",
    )}, sort_keys=True))


if __name__ == "__main__":
    main()
